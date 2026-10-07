"""Filter calibration report: what passed (possible false positives) and what was rejected but looks related
(possible misses), with the rule that decided every item. The user marks verdicts in Excel; the marked file is
the input for tuning config/filters/*.yaml."""

from __future__ import annotations

import asyncio
import random
import re
from collections import Counter
from collections.abc import Callable
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .client import NotFound, ProzorroClient, ProzorroError
from .db import Database
from .export import MONEY, _amount, _org, _Sheet
from .filter import RULE_UA, TenderFilter
from .probe import probe_hits, tender_digest
from .selection import TenderQuery, select_tenders
from .summary import tender_url
from .sync import parse_since

VERDICTS = '"так,ні,не знаю"'
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")


def cpv_digits_of(code: str | None) -> str:
    return (code or "").replace("-", "")


def _date_of_tender_id(tender_id: str | None) -> str | None:
    """UA-2026-07-15-001234-a -> 2026-07-15 (Prozorro ids carry the announcement date)."""
    m = re.match(r"UA-(\d{4}-\d{2}-\d{2})-", tender_id or "")
    return m.group(1) if m else None


def _risky(r: dict[str, Any], cpv_prefixes: tuple[str, ...]) -> bool:
    """A rejected tender with a stored digest that is worth re-checking."""
    if r["stage"] == "value" or r["probe"]:
        return True
    return any(c.replace("-", "").startswith(cpv_prefixes) for c in (r["cpvs"] or "").split(",") if c)


async def build_report(
    db: Database,
    client: ProzorroClient | None,
    f: TenderFilter,
    created_from: str | None = None,
    created_to: str | None = None,
    sample: int = 300,
    concurrency: int = 4,
    progress: Callable[[str], Any] | None = None,
    cpv_prefixes: list[str] | tuple[str, ...] = (),
    min_value: float | None = None,
    max_targeted: int = 1000,
    seed: int = 0,
) -> dict[str, Any]:
    """`sample` caps how many rejected tenders WITHOUT a stored digest (synced before digests existed) are
    re-fetched, in random order; their digests are saved, so repeated runs cover the whole period. Tenders WITH a
    digest are re-fetched only when risky: rejected by value, related words (probe), or a CPV from `cpv_prefixes`
    (up to `max_targeted`, largest first). `min_value` drops cheaper tenders (default: the filter threshold)."""
    progress = progress or (lambda m: None)
    prefixes = tuple(p.replace("-", "").strip() for p in cpv_prefixes if p.strip())
    floor = f.min_amount if min_value is None else min_value

    # 1. passed: every item of every relevant tender, with the rule that decided it
    passed_rows: list[dict[str, Any]] = []
    rules, topics, keywords, cpvs = Counter(), Counter(), Counter(), Counter()
    tenders = select_tenders(db, TenderQuery(profile=f.name, created_from=created_from, created_to=created_to))
    for t in tenders:
        d = f.evaluate(t)
        entity, _ = _org(t.get("procuringEntity"))
        for item, context in f.item_contexts(t):
            tr = f.trace_item(item, context)
            cls = item.get("classification") or {}
            rules[tr.rule] += 1
            if tr.match:
                topics[tr.match.topic] += 1
                cpvs[cls.get("id")] += 1
                if tr.keyword:
                    keywords[tr.keyword.lower()] += 1
            passed_rows.append(
                {
                    "tender": t,
                    "entity": entity,
                    "value": d.relevant_value,
                    "item": item.get("description"),
                    "cpv": cls.get("id"),
                    "cpv_name": cls.get("description"),
                    "matched": bool(tr.match),
                    "topic": tr.match.topic if tr.match else None,
                    "rule": RULE_UA.get(tr.rule, tr.rule),
                    "detail": tr.detail,
                }
            )
    progress(f"пройшли фільтр: {len(tenders)} тендерів, {len(passed_rows)} позицій")

    # 2. rejected candidates: re-fetch a sample of tenders rejected by topic or value and look for related words
    lo = parse_since(created_from).date().isoformat() if created_from else None
    hi = parse_since(created_to).date().isoformat() if created_to else None
    rejected = [
        r
        for r in db.rejected_decisions(f.name)
        if (d := _date_of_tender_id(r["tender_id"])) and (not lo or d >= lo) and (not hi or d < hi)
    ]
    known = [r for r in rejected if r["cpvs"] is not None]
    unknown = [r for r in rejected if r["cpvs"] is None]
    targeted = [r for r in known if r["stage"] == "value" or (_risky(r, prefixes) and (r["value"] or 0) >= floor)]
    targeted.sort(key=lambda r: -(r["value"] or 0))
    targeted = targeted[:max_targeted]
    random.Random(seed).shuffle(unknown)
    todo = targeted + unknown[:sample] if client else []
    candidates: list[dict[str, Any]] = []
    fetched = errors = backfilled = 0
    if todo:
        sem = asyncio.Semaphore(concurrency)

        async def check(r: dict[str, Any]) -> None:
            nonlocal fetched, errors, backfilled
            async with sem:
                try:
                    t = await client.get_tender(r["id"])
                except (NotFound, ProzorroError):
                    errors += 1
                    return
            fetched += 1
            if fetched % 100 == 0:
                progress(f"перевірено відкинутих: {fetched}/{len(todo)}")
            if r["cpvs"] is None:
                db.save_digest(r["id"], tender_digest(t))
                backfilled += 1
            d = f.evaluate(t)
            value = _amount(t.get("value"))
            if d.stage != "value" and (value or 0) < floor:
                return  # too cheap anyway: not a miss of the topic rules
            hits_title = probe_hits([t.get("title") or "", t.get("description") or ""])
            entity, _ = _org(t.get("procuringEntity"))
            for item, context in f.item_contexts(t):
                cls = item.get("classification") or {}
                hits = probe_hits([item.get("description") or ""])
                by_cpv = bool(prefixes) and cpv_digits_of(cls.get("id")).startswith(prefixes)
                if not (hits or hits_title or by_cpv or d.stage == "value"):
                    continue
                tr = f.trace_item(item, context)
                candidates.append(
                    {
                        "tender": t,
                        "entity": entity,
                        "value": value,
                        "relevant_value": d.relevant_value,
                        "stage": {"topic": "тема", "value": "вартість"}.get(d.stage, d.stage),
                        "reason": d.reason,
                        "item": item.get("description"),
                        "cpv": cls.get("id"),
                        "cpv_name": cls.get("description"),
                        "rule": RULE_UA.get(tr.rule, tr.rule),
                        "detail": tr.detail,
                        "hits": ", ".join(hits or hits_title) or ("код CPV" if by_cpv else ""),
                    }
                )

        progress(
            f"перевіряю {len(todo)} з {len(rejected)} відкинутих тендерів: {len(targeted)} ризикових "
            f"за збереженими даними, {len(todo) - len(targeted)} без збережених даних…"
        )
        await asyncio.gather(*(check(r) for r in todo))
    candidates.sort(key=lambda c: (-(c["value"] or 0), c["tender"].get("tenderID") or ""))

    return {
        "filter": f.name,
        "filter_key": f.key,
        "period": {"from": created_from, "to": created_to},
        "passed_tenders": len(tenders),
        "passed_items": passed_rows,
        "rules": rules,
        "topics": topics,
        "keywords": keywords,
        "cpvs": cpvs,
        "rejected_total": len(rejected),
        "rejected_with_digest": len(known),
        "rejected_targeted": len(targeted),
        "rejected_without_digest": len(unknown) - backfilled,
        "rejected_checked": fetched,
        "rejected_errors": errors,
        "candidates": candidates,
    }


def write_report(report: dict[str, Any], path: Path) -> dict[str, int]:
    wb = Workbook()
    ws = wb.active
    ws.title = "Як користуватися"
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 16
    lines = [
        (f"Калібрування фільтра «{report['filter']}»", Font(bold=True, size=14)),
        (f"Версія правил: {report['filter_key']}", None),
        (f"Період оголошення: {report['period']['from'] or '…'} — {report['period']['to'] or '…'}", None),
        ("", None),
        ("1. Аркуш «Пройшли»: позиції тендерів, що пройшли фільтр. У колонці «Вердикт» позначте «ні»,", None),
        ("   якщо тендер не за темою (хибне спрацювання). Позиції з «Збіг = ні» показують, що ще є в тендері.", None),
        ("2. Аркуш «Відкинуті кандидати»: тендери, які фільтр відкинув, але в яких є слова, схожі на", None),
        ("   кібербезпеку. Позначте «так», якщо тендер мав пройти (пропуск).", None),
        ("3. Надішліть файл з позначками: за ним будуть уточнені коди, ключові слова та виключення.", None),
    ]
    for text, font in lines:
        ws.append([text])
        if font:
            ws.cell(row=ws.max_row, column=1).font = font
    ws.append([])
    for title, counter in [
        ("Правила, за якими вирішено позиції (тендери, що пройшли)", report["rules"]),
        ("Теми збігів", report["topics"]),
        ("Ключові слова, що спрацювали", report["keywords"]),
        ("Коди CPV позицій, що пройшли", report["cpvs"]),
    ]:
        ws.append([title])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        for k, v in counter.most_common(30):
            ws.append([RULE_UA.get(k, k) if counter is report["rules"] else k, v])
        ws.append([])
    ws.append(["Відкинуті тендери за період", report["rejected_total"]])
    ws.append(["з них зі збереженими даними (тема, коди, слова)", report["rejected_with_digest"]])
    ws.append(["  ризикових серед них (перевірено повторно)", report["rejected_targeted"]])
    ws.append(["ще без збережених даних (запустіть калібрування ще раз)", report["rejected_without_digest"]])
    ws.append(["перевірено повторно за цей запуск", report["rejected_checked"]])
    ws.append(["помилок завантаження", report["rejected_errors"]])
    ws.append(["позицій-кандидатів на пропуск", len(report["candidates"])])

    passed = _Sheet(
        wb.create_sheet("Пройшли"),
        [
            ("Тендер", 24, None),
            ("Назва", 40, None),
            ("Замовник", 30, None),
            ("Вартість релевантних лотів", 16, MONEY),
            ("Позиція", 50, None),
            ("CPV", 13, None),
            ("Назва CPV", 30, None),
            ("Збіг", 7, None),
            ("Тема", 16, None),
            ("Правило", 26, None),
            ("Деталі", 36, None),
            ("Вердикт", 10, None),
            ("Коментар", 30, None),
        ],
    )
    for r in report["passed_items"]:
        t = r["tender"]
        passed.add(
            [
                t.get("tenderID"),
                t.get("title"),
                r["entity"],
                r["value"],
                r["item"],
                r["cpv"],
                r["cpv_name"],
                "так" if r["matched"] else "ні",
                r["topic"],
                r["rule"],
                r["detail"],
                None,
                None,
            ],
            link=tender_url(t),
        )
    cand = _Sheet(
        wb.create_sheet("Відкинуті кандидати"),
        [
            ("Тендер", 24, None),
            ("Назва", 40, None),
            ("Замовник", 30, None),
            ("Очікувана вартість", 16, MONEY),
            ("Вартість релевантних лотів", 16, MONEY),
            ("Відкинуто на етапі", 12, None),
            ("Причина", 34, None),
            ("Позиція", 50, None),
            ("CPV", 13, None),
            ("Назва CPV", 30, None),
            ("Правило", 26, None),
            ("Деталі", 36, None),
            ("Схожі слова", 24, None),
            ("Мав пройти?", 11, None),
            ("Коментар", 30, None),
        ],
    )
    for r in report["candidates"]:
        t = r["tender"]
        cand.add(
            [
                t.get("tenderID"),
                t.get("title"),
                r["entity"],
                r["value"],
                r["relevant_value"],
                r["stage"],
                r["reason"],
                r["item"],
                r["cpv"],
                r["cpv_name"],
                r["rule"],
                r["detail"],
                r["hits"],
                None,
                None,
            ],
            link=tender_url(t),
        )
    for sheet, col in ((passed, "L"), (cand, "N")):
        dv = DataValidation(type="list", formula1=VERDICTS, allow_blank=True)
        sheet.ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{max(sheet.ws.max_row, 2)}")
        for row in range(2, sheet.ws.max_row + 1):
            sheet.ws[f"{col}{row}"].fill = NOTE_FILL
    counts = {"Пройшли": passed.finish(), "Відкинуті кандидати": cand.finish()}
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return counts
