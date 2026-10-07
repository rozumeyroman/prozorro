"""Excel export of stored tenders: tenders, items, winners, bids and unit prices on separate sheets."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .filter import TenderFilter
from .settings import KYIV_TZ
from .summary import tender_url

MONEY = "#,##0.00"
DATETIME = "dd.mm.yyyy hh:mm"
PERCENT = "0.0%"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
LINK_FONT = Font(color="0563C1", underline="single")

STATUS_UA = {
    "active.enquiries": "Період уточнень",
    "active.tendering": "Прийом пропозицій",
    "active.pre-qualification": "Прекваліфікація",
    "active.pre-qualification.stand-still": "Прекваліфікація (оскарження)",
    "active.auction": "Аукціон",
    "active.qualification": "Кваліфікація",
    "active.qualification.stand-still": "Кваліфікація (оскарження)",
    "active.awarded": "Переможця визначено",
    "complete": "Завершено",
    "cancelled": "Скасовано",
    "unsuccessful": "Не відбувся",
}
AWARD_STATUS_UA = {
    "active": "Переможець",
    "pending": "Розглядається",
    "unsuccessful": "Відхилено",
    "cancelled": "Скасовано",
}
CONTRACT_STATUS_UA = {
    "active": "Підписано",
    "pending": "Очікує підписання",
    "terminated": "Виконано/розірвано",
    "cancelled": "Скасовано",
}

# (header, width, number format)
Column = tuple[str, int, str | None]


def _dt(value: str | None) -> datetime | None:
    """ISO timestamp -> naive Kyiv time (Excel has no time zones)."""
    if not value:
        return None
    return datetime.fromisoformat(value).astimezone(KYIV_TZ).replace(tzinfo=None)


def _amount(v: dict[str, Any] | None) -> float | None:
    return (v or {}).get("amount")


def _vat(v: dict[str, Any] | None) -> str | None:
    if not v or v.get("valueAddedTaxIncluded") is None:
        return None
    return "з ПДВ" if v["valueAddedTaxIncluded"] else "без ПДВ"


def _org(org: dict[str, Any] | None) -> tuple[str | None, str | None]:
    org = org or {}
    ident = org.get("identifier") or {}
    return org.get("name") or ident.get("legalName"), ident.get("id")


class _Sheet:
    def __init__(self, ws: Worksheet, columns: list[Column]):
        self.ws = ws
        self.columns = columns
        ws.append([c[0] for c in columns])
        for i, (_, width, _) in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=i)
            cell.fill, cell.font = HEADER_FILL, HEADER_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "B2"
        # Printing: landscape, all columns on one page width.
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0

    def add(self, values: list[Any], link: str | None = None) -> None:
        self.ws.append(values)
        row = self.ws.max_row
        for i, (_, _, fmt) in enumerate(self.columns, start=1):
            if fmt:
                self.ws.cell(row=row, column=i).number_format = fmt
        if link:
            cell = self.ws.cell(row=row, column=1)
            cell.hyperlink, cell.font = link, LINK_FONT

    def finish(self) -> int:
        if self.ws.max_row > 1:
            self.ws.auto_filter.ref = self.ws.dimensions
        return self.ws.max_row - 1


def export_tenders(
    tenders: list[dict[str, Any]],
    path: Path,
    tender_filter: TenderFilter,
    include_summary: bool = True,
    period_note: str | None = None,
) -> dict[str, int]:
    wb = Workbook()
    s_tenders = _Sheet(
        wb.active,
        [
            ("Тендер", 24, None),
            ("Назва", 50, None),
            ("Замовник", 40, None),
            ("ЄДРПОУ замовника", 14, None),
            ("Регіон", 20, None),
            ("Тема", 18, None),
            ("Очікувана вартість", 16, MONEY),
            ("Вартість релевантних лотів", 16, MONEY),
            ("Валюта", 8, None),
            ("ПДВ", 9, None),
            ("Статус", 20, None),
            ("Тип процедури", 18, None),
            ("Створено", 16, DATETIME),
            ("Пропозиції до", 16, DATETIME),
            ("Учасників", 10, None),
            ("Переможець", 36, None),
            ("ЄДРПОУ переможця", 14, None),
            ("Сума переможця", 16, MONEY),
            ("Дата рішення", 16, DATETIME),
            ("Знижка від очікуваної", 12, PERCENT),
            ("Договір", 16, None),
            ("Дата договору", 16, DATETIME),
            ("Сума договору", 16, MONEY),
            ("Документів", 11, None),
        ],
    )
    wb.active.title = "Тендери"
    s_items = _Sheet(
        wb.create_sheet("Позиції"),
        [
            ("Тендер", 24, None),
            ("Лот", 30, None),
            ("Позиція", 60, None),
            ("CPV", 13, None),
            ("CPV назва", 40, None),
            ("Кількість", 10, "#,##0.###"),
            ("Одиниця", 12, None),
            ("Відповідає темі", 14, None),
            ("Причина збігу", 36, None),
        ],
    )
    s_awards = _Sheet(
        wb.create_sheet("Переможці"),
        [
            ("Тендер", 24, None),
            ("Назва", 40, None),
            ("Замовник", 36, None),
            ("Лот", 30, None),
            ("Статус рішення", 14, None),
            ("Учасник", 40, None),
            ("ЄДРПОУ", 14, None),
            ("Сума", 16, MONEY),
            ("ПДВ", 9, None),
            ("Дата рішення", 16, DATETIME),
            ("Очікувана вартість", 16, MONEY),
            ("Знижка", 10, PERCENT),
            ("Учасників", 10, None),
            ("Статус договору", 16, None),
            ("Дата договору", 16, DATETIME),
            ("Сума договору", 16, MONEY),
        ],
    )
    s_bids = _Sheet(
        wb.create_sheet("Пропозиції"),
        [
            ("Тендер", 24, None),
            ("Лот", 30, None),
            ("Учасник", 40, None),
            ("ЄДРПОУ", 14, None),
            ("Сума", 16, MONEY),
            ("Початкова сума", 16, MONEY),
            ("ПДВ", 9, None),
            ("Статус", 12, None),
            ("Переможець", 11, None),
            ("Дата", 16, DATETIME),
        ],
    )
    s_prices = _Sheet(
        wb.create_sheet("Ціни за одиницю"),
        [
            ("Тендер", 24, None),
            ("Учасник", 40, None),
            ("ЄДРПОУ", 14, None),
            ("Позиція", 60, None),
            ("CPV", 13, None),
            ("Кількість", 10, "#,##0.###"),
            ("Одиниця", 12, None),
            ("Ціна за одиницю", 16, MONEY),
            ("ПДВ", 9, None),
            ("Переможець", 11, None),
        ],
    )

    for t in tenders:
        _add_tender(t, tender_filter, s_tenders, s_items, s_awards, s_bids, s_prices)

    counts = {}
    if include_summary:
        from .analytics import summarize  # analytics reuses this module's helpers

        ws = wb.create_sheet("Аналітика", 0)
        counts["Аналітика"] = write_summary(ws, summarize(tenders, tender_filter), period_note)
        wb.active = 0
    counts |= {
        "Тендери": s_tenders.finish(),
        "Позиції": s_items.finish(),
        "Переможці": s_awards.finish(),
        "Пропозиції": s_bids.finish(),
        "Ціни за одиницю": s_prices.finish(),
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return counts


def _add_tender(t, tender_filter, s_tenders, s_items, s_awards, s_bids, s_prices) -> None:
    tid, url = t.get("tenderID"), tender_url(t)
    entity, entity_code = _org(t.get("procuringEntity"))
    decision = tender_filter.evaluate(t)
    matched = {m.item_id: m for m in decision.matches}
    lots = {lot["id"]: lot for lot in t.get("lots") or [] if lot.get("id")}
    bids = t.get("bids") or []
    awards = t.get("awards") or []
    contracts = t.get("contracts") or []
    contracts_by_award = {c.get("awardID"): c for c in contracts if c.get("awardID")}
    winners = [a for a in awards if a.get("status") == "active"]
    if len(winners) == 1 and len(contracts) == 1 and not contracts_by_award:
        # Contract without awardID (old data): with a single winner the match is unambiguous.
        contracts_by_award[winners[0].get("id")] = contracts[0]
    winning_bids = {a.get("bid_id") for a in winners}
    value = t.get("value") or {}

    # Tenders: one row; winner columns are filled when there is exactly one winner (multi-lot -> see sheet).
    w = winners[0] if len(winners) == 1 else None
    wc = contracts_by_award.get(w.get("id")) if w else None
    w_name, w_code = _org((w.get("suppliers") or [None])[0]) if w else (None, None)
    if len(winners) > 1:
        w_name = f"{len(winners)} переможців (див. аркуш «Переможці»)"
    expected = _expected_for_award(w, t, lots) if w else None
    s_tenders.add(
        [
            tid,
            t.get("title"),
            entity,
            entity_code,
            ((t.get("procuringEntity") or {}).get("address") or {}).get("region"),
            ", ".join(decision.topics),
            _amount(value),
            decision.relevant_value,
            value.get("currency"),
            _vat(value),
            STATUS_UA.get(t.get("status"), t.get("status")),
            t.get("procurementMethodType"),
            _dt(t.get("dateCreated") or t.get("date")),
            _dt((t.get("tenderPeriod") or {}).get("endDate")),
            len([b for b in bids if b.get("status") in ("active", "pending", None)]) or None,
            w_name,
            w_code,
            _amount(w.get("value")) if w else None,
            _dt(w.get("date")) if w else None,
            _discount(w.get("value") if w else None, expected),
            CONTRACT_STATUS_UA.get(wc.get("status"), wc.get("status")) if wc else None,
            _dt(wc.get("dateSigned")) if wc else None,
            _amount(wc.get("value")) if wc else None,
            len({d.get("id") for d in t.get("documents") or []}),
        ],
        link=url,
    )

    for it in t.get("items") or []:
        cls = it.get("classification") or {}
        m = matched.get(it.get("id", ""))
        s_items.add(
            [
                tid,
                (lots.get(it.get("relatedLot") or "") or {}).get("title"),
                it.get("description"),
                cls.get("id"),
                cls.get("description"),
                it.get("quantity"),
                (it.get("unit") or {}).get("name"),
                "так" if m else "ні",
                m.reason if m else None,
            ],
            link=url,
        )

    for a in awards:
        name, code = _org((a.get("suppliers") or [None])[0])
        c = contracts_by_award.get(a.get("id"))
        exp = _expected_for_award(a, t, lots)
        lot_id = a.get("lotID")
        n_bids = len(
            [b for b in bids if not lot_id or any(lv.get("relatedLot") == lot_id for lv in b.get("lotValues") or [])]
        )
        s_awards.add(
            [
                tid,
                t.get("title"),
                entity,
                (lots.get(lot_id or "") or {}).get("title"),
                AWARD_STATUS_UA.get(a.get("status"), a.get("status")),
                name,
                code,
                _amount(a.get("value")),
                _vat(a.get("value")),
                _dt(a.get("date")),
                _amount(exp),
                _discount(a.get("value"), exp),
                n_bids or None,
                CONTRACT_STATUS_UA.get(c.get("status"), c.get("status")) if c else None,
                _dt(c.get("dateSigned")) if c else None,
                _amount(c.get("value")) if c else None,
            ],
            link=url,
        )

    for b in bids:
        name, code = _org((b.get("tenderers") or [None])[0])
        is_winner = "так" if b.get("id") in winning_bids else None
        entries = [(lv.get("relatedLot"), lv.get("value"), lv.get("initialValue")) for lv in b.get("lotValues") or []]
        if not entries:
            entries = [(None, b.get("value"), b.get("initialValue"))]
        for lot_id, v, initial in entries:
            s_bids.add(
                [
                    tid,
                    (lots.get(lot_id or "") or {}).get("title"),
                    name,
                    code,
                    _amount(v),
                    _amount(initial),
                    _vat(v),
                    b.get("status"),
                    is_winner,
                    _dt(b.get("date")),
                ],
                link=url,
            )
        for it in b.get("items") or []:
            unit = it.get("unit") or {}
            if _amount(unit.get("value")) is None:
                continue
            s_prices.add(
                [
                    tid,
                    name,
                    code,
                    it.get("description"),
                    (it.get("classification") or {}).get("id"),
                    it.get("quantity"),
                    unit.get("name"),
                    _amount(unit.get("value")),
                    _vat(unit.get("value")),
                    is_winner,
                ],
                link=url,
            )


def _expected_for_award(
    award: dict[str, Any] | None, tender: dict[str, Any], lots: dict[str, Any]
) -> dict[str, Any] | None:
    """Expected value of the lot the award belongs to (or of the whole tender)."""
    if not award:
        return None
    lot = lots.get(award.get("lotID") or "")
    return (lot or tender).get("value")


def _discount(price: dict[str, Any] | None, expected: dict[str, Any] | None) -> float | None:
    """1 - price / expected, only when both amounts are comparable (same currency and VAT basis)."""
    p, e = _amount(price), _amount(expected)
    if p is None or not e:
        return None
    if (price or {}).get("currency") != (expected or {}).get("currency"):
        return None
    if (price or {}).get("valueAddedTaxIncluded") != (expected or {}).get("valueAddedTaxIncluded"):
        return None
    return round(1 - p / e, 4)


SECTION_FONT = Font(bold=True, size=12, color="1F4E78")


def write_summary(ws: Worksheet, summary: dict[str, Any], period_note: str | None = None) -> int:
    """Summary sheet: key figures and breakdown tables one under another. Returns the number of filled rows."""
    ws.column_dimensions["A"].width = 44
    for col in "BCDE":
        ws.column_dimensions[col].width = 18

    def section(title: str) -> None:
        if ws.max_row > 1:
            ws.append([])
        ws.append([title])
        ws.cell(row=ws.max_row, column=1).font = SECTION_FONT

    def header(*names: str) -> None:
        ws.append(list(names))
        for i in range(1, len(names) + 1):
            c = ws.cell(row=ws.max_row, column=i)
            c.fill, c.font = HEADER_FILL, HEADER_FONT

    def row(values: list[Any], formats: list[str | None]) -> None:
        ws.append(values)
        for i, fmt in enumerate(formats, start=1):
            if fmt:
                ws.cell(row=ws.max_row, column=i).number_format = fmt

    ws.append([f"Аналітика за фільтром «{summary['filter']}»"])
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    if period_note:
        ws.append([period_note])

    section("Підсумок")
    comp, disc = summary["competition"], summary["discount"]
    for label, value, fmt in [
        ("Тендерів", summary["tenders"], None),
        ("Очікувана вартість релевантних лотів, грн", summary["expected_value_uah"], MONEY),
        ("Рішень про переможця", summary["awards"], None),
        ("Сума рішень про переможця, грн", summary["awarded_value_uah"], MONEY),
        ("Середня кількість учасників", comp.get("avg_bidders"), "0.00"),
        ("Частка торгів з одним учасником", comp.get("single_bidder_share"), PERCENT),
        ("Медіанна знижка від очікуваної вартості", disc.get("median"), PERCENT),
        ("Середня знижка від очікуваної вартості", disc.get("mean"), PERCENT),
    ]:
        row([label, value], [None, fmt])
    if summary["non_uah_tenders"]:
        ws.append([f"Тендерів в іншій валюті (не враховані в сумах): {summary['non_uah_tenders']}"])

    for title, key in [("За статусом", "by_status"), ("За темою", "by_topic")]:
        section(title)
        header("", "Тендерів", "Очікувана вартість, грн")
        for name, v in sorted(summary[key].items(), key=lambda kv: -kv[1]["expected"]):
            row([name, v["count"], v["expected"]], [None, None, MONEY])

    section("За місяцем оголошення")
    header("Місяць", "Тендерів", "Очікувана вартість, грн", "Сума рішень, грн")
    for month, v in summary["by_month"].items():
        row([month, v["count"], v["expected"], v["awarded"]], [None, None, MONEY, MONEY])

    for title, key, amount_label in [
        ("Топ переможців за сумою", "top_winners_by_amount", "Сума рішень, грн"),
        ("Топ переможців за кількістю перемог", "top_winners_by_count", "Сума рішень, грн"),
        ("Топ замовників за сумою", "top_buyers_by_amount", "Очікувана вартість, грн"),
        ("Топ замовників за кількістю тендерів", "top_buyers_by_count", "Очікувана вартість, грн"),
    ]:
        section(title)
        header("Назва", "ЄДРПОУ", "Кількість", amount_label)
        for r in summary[key]:
            row([r["name"], r["edrpou"], r["count"], r["amount"]], [None, None, None, MONEY])
    return ws.max_row
