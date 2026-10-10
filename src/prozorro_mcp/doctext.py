"""Plain text from downloaded tender documents: PDF, DOCX, XLSX, ODT/ODS, text formats, ZIP archives and files
signed with a qualified e-signature (.p7s/.p7m with the document inside). Scanned PDFs have no text layer: they
are reported as such; ocr_file() recognises them with tesseract on request (offers prepare/text --ocr)."""

from __future__ import annotations

import io
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from html import unescape
from pathlib import Path
from xml.etree import ElementTree

MAX_DEPTH = 3  # signed file inside an archive inside a signed file is the deepest seen in practice
MAX_ARCHIVE_MEMBERS = 50
SCAN_CHARS_PER_PAGE = 30  # less text than this per PDF page means the PDF is (mostly) a scan

TEXT_SUFFIXES = {".txt", ".csv", ".json", ".xml", ".html", ".htm", ".md"}


@dataclass
class Extracted:
    text: str
    note: str | None = None  # why there is little or no text: scan, unsupported format, error

    @property
    def ok(self) -> bool:
        return bool(self.text.strip())


def extract_text(path: Path) -> Extracted:
    try:
        return _extract(path.read_bytes(), path.name, 0)
    except OSError as e:
        return Extracted("", f"не вдалося прочитати файл: {e}")


def _kind(data: bytes, name: str) -> str:
    """Detect by content first: Prozorro file titles often lie about the format (or have none)."""
    head = data[:8]
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"PK\x03\x04"):
        try:
            names = set(zipfile.ZipFile(io.BytesIO(data)).namelist())
        except zipfile.BadZipFile:
            return "unknown"
        if "word/document.xml" in names:
            return "docx"
        if "xl/workbook.xml" in names:
            return "xlsx"
        if "content.xml" in names and "mimetype" in names:
            return "odf"
        return "zip"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "ole"  # old .doc/.xls
    if head.startswith(b"Rar!") or head.startswith(b"7z\xbc\xaf"):
        return "rar7z"
    if head[:1] == b"\x30":  # DER SEQUENCE: CMS (signed data)
        return "cms"
    suffix = Path(name).suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return "text"
    if head[:4] in (b"\x89PNG", b"\xff\xd8\xff\xe0", b"\xff\xd8\xff\xe1", b"\xff\xd8\xff\xdb") or suffix in {
        ".jpg",
        ".jpeg",
        ".png",
        ".tif",
        ".tiff",
    }:
        return "image"
    try:
        data[:4096].decode("utf-8")
        return "text"
    except UnicodeDecodeError:
        return "unknown"


def _extract(data: bytes, name: str, depth: int) -> Extracted:
    kind = _kind(data, name)
    try:
        if kind == "pdf":
            return _pdf(data)
        if kind == "docx":
            return Extracted(_docx(data))
        if kind == "xlsx":
            return Extracted(_xlsx(data))
        if kind == "odf":
            return Extracted(_odf(data))
        if kind == "text":
            return Extracted(_text(data, name))
        if kind == "cms":
            return _cms(data, name, depth)
        if kind == "zip":
            return _zip(data, depth)
    except Exception as e:  # noqa: BLE001 - a broken file must not stop a batch
        return Extracted("", f"помилка читання ({kind}): {e}")
    notes = {
        "ole": "старий формат Word/Excel (.doc/.xls): відкрийте файл вручну",
        "rar7z": "архів RAR/7z: розпакуйте вручну",
        "image": "зображення (скан): потрібне розпізнавання тексту (OCR)",
    }
    return Extracted("", notes.get(kind, "невідомий формат"))


def _pdf(data: bytes) -> Extracted:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:  # noqa: BLE001
            text = ""
        pages.append(f"[сторінка {i}]\n{text.strip()}")
    text = "\n\n".join(pages)
    chars = sum(len(p) for p in pages) - sum(len(f"[сторінка {i}]\n") for i in range(1, len(pages) + 1))
    if chars < SCAN_CHARS_PER_PAGE * max(1, len(pages)):
        return Extracted(text if chars else "", f"PDF без текстового шару (скан, {len(pages)} с.): потрібне OCR")
    return Extracted(text)


W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx(data: bytes) -> str:
    z = zipfile.ZipFile(io.BytesIO(data))
    body = ElementTree.fromstring(z.read("word/document.xml")).find(f"{W_NS}body")

    def text_of(el: ElementTree.Element) -> str:
        return "".join(t.text or "" for t in el.iter(f"{W_NS}t"))

    lines = []
    for block in body if body is not None else []:
        if block.tag == f"{W_NS}p":
            lines.append(text_of(block))
        elif block.tag == f"{W_NS}tbl":
            for tr in block.iter(f"{W_NS}tr"):
                lines.append(" | ".join(text_of(tc).strip() for tc in tr.findall(f"{W_NS}tc")))
    return "\n".join(line for line in lines if line.strip())


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"[аркуш {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                out.append(" | ".join(cells).rstrip(" |"))
    return "\n".join(out)


def _odf(data: bytes) -> str:
    z = zipfile.ZipFile(io.BytesIO(data))
    root = ElementTree.fromstring(z.read("content.xml"))
    text_ns = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
    return "\n".join("".join(p.itertext()) for p in root.iter() if p.tag in (f"{text_ns}p", f"{text_ns}h"))


def _text(data: bytes, name: str) -> str:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = data.decode("utf-8", errors="replace")
    if Path(name).suffix.lower() in {".html", ".htm", ".xml"}:
        text = unescape(re.sub(r"<[^>]+>", " ", text))
        text = re.sub(r"[ \t]+", " ", text)
    return text


def _cms(data: bytes, name: str, depth: int) -> Extracted:
    """A file signed with an attached signature: the document is inside. A detached signature has no content."""
    from asn1crypto import cms

    info = cms.ContentInfo.load(data)
    if info["content_type"].native != "signed_data":
        return Extracted("", "підписаний файл незнайомого типу")
    content = info["content"]["encap_content_info"]["content"]
    inner = content.native if content is not None else None
    if not inner:
        return Extracted("", "окремий файл підпису (без документа всередині)")
    if depth >= MAX_DEPTH:
        return Extracted("", "забагато вкладених рівнів")
    inner_name = re.sub(r"\.(p7s|p7m|sig)$", "", name, flags=re.I)
    return _extract(bytes(inner), inner_name, depth + 1)


def _zip(data: bytes, depth: int) -> Extracted:
    if depth >= MAX_DEPTH:
        return Extracted("", "забагато вкладених архівів")
    z = zipfile.ZipFile(io.BytesIO(data))
    parts, notes = [], []
    members = [m for m in z.infolist() if not m.is_dir()][:MAX_ARCHIVE_MEMBERS]
    for m in members:
        if m.file_size > 50 * 1_048_576:
            notes.append(f"{m.filename}: завеликий")
            continue
        r = _extract(z.read(m), m.filename, depth + 1)
        if r.ok:
            parts.append(f"===== {m.filename} =====\n{r.text}")
        if r.note:
            notes.append(f"{m.filename}: {r.note}")
    return Extracted("\n\n".join(parts), "; ".join(notes) or None)


OCR_HINT = "для OCR встановіть tesseract і poppler (macOS: brew install tesseract poppler)"


def ocr_missing_tools() -> list[str]:
    return [t for t in ("tesseract", "pdftoppm") if not shutil.which(t)]


def _signed_inner(data: bytes) -> bytes | None:
    from asn1crypto import cms

    try:
        info = cms.ContentInfo.load(data)
        content = info["content"]["encap_content_info"]["content"]
        inner = content.native if content is not None else None
        return bytes(inner) if inner else None
    except Exception:  # noqa: BLE001
        return None


def ocr_file(path: Path, pages: int = 2, lang: str = "eng", timeout: int = 180) -> Extracted:
    """Recognise the first `pages` pages of a scanned PDF (or an image) with tesseract. Only for scans: costly."""
    missing = ocr_missing_tools()
    if "tesseract" in missing:
        return Extracted("", f"OCR недоступне: {OCR_HINT}")
    try:
        data = path.read_bytes()
    except OSError as e:
        return Extracted("", f"не вдалося прочитати файл: {e}")
    kind = _kind(data, path.name)
    if kind == "cms":
        inner = _signed_inner(data)
        if not inner:
            return Extracted("", "OCR: підписаний файл без документа всередині")
        data = inner
        kind = _kind(data, re.sub(r"\.(p7s|p7m)$", "", path.name, flags=re.I))
    with tempfile.TemporaryDirectory(prefix="prozorro-ocr-") as tmp_dir:
        tmp = Path(tmp_dir)
        if kind == "pdf":
            if "pdftoppm" in missing:
                return Extracted("", f"OCR PDF недоступне: {OCR_HINT}")
            src = tmp / "in.pdf"
            src.write_bytes(data)
            try:
                subprocess.run(
                    ["pdftoppm", "-f", "1", "-l", str(pages), "-r", "200", "-png", str(src), str(tmp / "p")],
                    check=True,
                    capture_output=True,
                    timeout=timeout,
                )
            except (subprocess.SubprocessError, OSError) as e:
                return Extracted("", f"OCR: не вдалося перетворити PDF на зображення ({e})")
            images = sorted(tmp.glob("p*.png"))
        elif kind == "image":
            src = tmp / ("in" + (Path(path.name).suffix or ".png"))
            src.write_bytes(data)
            images = [src]
        else:
            return Extracted("", f"OCR: формат {kind} не підтримується")
        texts = []
        for i, img in enumerate(images, start=1):
            try:
                r = subprocess.run(
                    ["tesseract", str(img), "-", "-l", lang],
                    capture_output=True,
                    text=True,
                    timeout=timeout,
                )
            except (subprocess.SubprocessError, OSError) as e:
                return Extracted("", f"OCR: помилка tesseract ({e})")
            texts.append(f"[сторінка {i}, OCR]\n{r.stdout.strip()}")
    text = "\n\n".join(texts)
    return Extracted(text, None if text.strip() else "OCR не знайшло тексту")
