from openpyxl import load_workbook
from test_sync import run_sync

from prozorro_mcp.calibrate import build_report, write_report
from prozorro_mcp.client import ProzorroClient
from prozorro_mcp.db import Database
from prozorro_mcp.settings import Settings


async def test_calibration_report(fake, cyber_filter, tmp_path):
    db = Database(":memory:")
    await run_sync(fake, cyber_filter, db)
    async with ProzorroClient(Settings(api_url="http://fake/api/2.5", max_retries=0), transport=fake.transport()) as c:
        report = await build_report(db, c, cyber_filter, sample=50)

    assert report["passed_tenders"] == 1
    assert [r["item"] for r in report["passed_items"]] == ["Ліцензії ESET PROTECT Complete, 300 робочих місць"]
    assert report["rules"]["strong"] == 1
    cand = {(r["tender"]["title"], r["stage"]) for r in report["candidates"]}
    # firewall licence under the threshold, and Cisco switches that the cyber profile does not take
    assert ("Міжмережевий екран", "вартість") in cand
    assert ("Закупівля комутаторів для ЦОД", "тема") in cand
    # cheap tenders rejected by topic are not candidates
    assert not any(t == "Закупівля картриджів" for t, _ in cand)

    path = tmp_path / "cal.xlsx"
    rows = write_report(report, path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Як користуватися", "Пройшли", "Відкинуті кандидати"]
    assert rows["Пройшли"] == 1 and rows["Відкинуті кандидати"] == len(report["candidates"])
    header = [c.value for c in wb["Відкинуті кандидати"][1]]
    assert "Мав пройти?" in header and wb["Відкинуті кандидати"].data_validations.dataValidation


async def test_calibration_offline(fake, cyber_filter):
    db = Database(":memory:")
    await run_sync(fake, cyber_filter, db)
    report = await build_report(db, None, cyber_filter)
    assert report["passed_tenders"] == 1 and report["rejected_checked"] == 0 and report["rejected_total"] > 0
