"""Coverage check against the prozorro.gov.ua search and finding tenders by UA-… id."""

import json
from datetime import timedelta

import pytest
from conftest import NOW
from openpyxl import load_workbook
from test_sync import run_sync

from prozorro_mcp import cli
from prozorro_mcp.client import ProzorroClient
from prozorro_mcp.coverage import check_coverage
from prozorro_mcp.db import Database
from prozorro_mcp.settings import Settings
from prozorro_mcp.site import SiteClient, resolve_internal_id

DAY = NOW.strftime("%Y-%m-%d")


def settings(**kw):
    return Settings(
        api_url="http://fake/api/2.5",
        site_url="http://fake",
        max_retries=0,
        feed_retry_delay=0,
        site_min_interval=0,
        **kw,
    )


async def coverage(fake, f, db, **kw):
    s = settings()
    async with ProzorroClient(s, transport=fake.transport()) as c, SiteClient(s, transport=fake.transport()) as site:
        return await check_coverage(db, site, f, date_from=DAY, date_to=DAY, client=c, **kw)


async def test_coverage_explains_every_tender(fake, tender_filter, tenders):
    db = Database(":memory:")
    await run_sync(fake, tender_filter, db)
    excluded = next(t for t in tenders if t["title"] == "Антивірусний захист")
    db.add_exclusion(excluded["tenderID"], "дубль")
    # a tender the feed walk never saw
    missed = next(t for t in tenders if t["title"] == "Серверне обладнання")
    db.conn.execute("DELETE FROM tenders WHERE id = ?", (missed["id"],))
    db.conn.execute("DELETE FROM filter_decisions WHERE id = ?", (missed["id"],))
    db.commit()

    cpvs = ["32420000-3", "30230000-0", "48760000-3", "48730000-4", "39100000-3"]
    report = await coverage(fake, tender_filter, db, cpvs=cpvs, min_value=500_000)
    by_id = {r["tenderID"]: r for r in report["not_in_db"]}
    assert by_id[excluded["tenderID"]]["reason"] == "виключено вручну" and by_id[excluded["tenderID"]]["detail"]
    assert by_id[missed["tenderID"]]["reason"] == "немає в базі"
    furniture = next(t for t in tenders if t["title"] == "Меблі та мережеве обладнання")
    assert by_id[furniture["tenderID"]]["reason"] == "відкинуто фільтром"
    assert "value" in by_id[furniture["tenderID"]]["detail"]
    # 400k firewall tender is below the value threshold: not counted at all
    firewall = next(t for t in tenders if t["title"] == "Міжмережевий екран")
    assert firewall["tenderID"] not in by_id
    # created yesterday (before the synced period) but accepting bids today: missing, with a hint why
    old = next(t for t in tenders if t["title"] == "Старий тендер")
    assert by_id[old["tenderID"]]["reason"] == "немає в базі" and "до періоду" in by_id[old["tenderID"]]["detail"]
    assert report["missing"] == 2 and report["missing_pct"] > 0

    # fetch the missing one by id: found via the tender page, fetched and stored
    report = await coverage(fake, tender_filter, db, cpvs=cpvs, min_value=500_000, fetch=True)
    assert report["missing"] == 0
    assert db.get_tender(missed["tenderID"])["title"] == "Серверне обладнання"
    assert any(p == f"/api/tenders/{missed['tenderID']}/summary" for p in fake.requests)


async def test_resolve_internal_id_checks_candidates(fake, tenders):
    db = Database(":memory:")
    t = tenders[0]
    s = settings()
    async with ProzorroClient(s, transport=fake.transport()) as c, SiteClient(s, transport=fake.transport()) as site:
        assert await resolve_internal_id(t["tenderID"], db, c, site) == t["id"]
        assert await resolve_internal_id("UA-2000-01-01-000000-a", db, c, site) is None


def run(capsys, *argv):
    cli.main(list(argv))
    return json.loads(capsys.readouterr().out)


@pytest.fixture
def served(tmp_path, monkeypatch, fake):
    server, api_url = fake.serve()
    monkeypatch.setenv("PROZORRO_API_URL", api_url)
    monkeypatch.setenv("PROZORRO_SITE_URL", api_url.rsplit("/api/", 1)[0])
    monkeypatch.setenv("PROZORRO_DB", str(tmp_path / "db.sqlite"))
    monkeypatch.setenv("PROZORRO_OUTPUT_DIR", str(tmp_path / "out"))
    monkeypatch.setenv("PROZORRO_FILTER", "it-infrastructure")
    monkeypatch.setenv("PROZORRO_FEED_RETRY_DELAY", "0")
    monkeypatch.setenv("PROZORRO_SITE_INTERVAL", "0")
    yield fake
    server.shutdown()


def test_cli_exclude_coverage_export(served, tmp_path, capsys, tenders):
    fake = served
    # the oldest demo tender (yesterday 07:00) is before `since`: the walk ends on it
    since = (NOW.replace(hour=8, minute=0) - timedelta(days=1)).isoformat()
    stats = run(capsys, "sync", "--since", since, "--shards", "2")
    assert stats["complete"] is True and len(stats["shard_details"]) == 2
    relevant = stats["relevant_found"]
    t = next(t for t in fake.tenders.values() if t["title"] == "Антивірусний захист")
    assert run(capsys, "exclude", "add", t["tenderID"], "--reason", "не той профіль")[0]["title"]
    assert run(capsys, "exclude", "list")[0]["reason"] == "не той профіль"
    out = run(capsys, "export", "-o", str(tmp_path / "x.xlsx"))
    assert out["rows"]["Виключені"] == 1 and out["rows"]["Тендери"] == relevant - 1
    ws = load_workbook(tmp_path / "x.xlsx")["Виключені"]
    assert ws["A2"].value == t["tenderID"] and ws["C2"].value == "не той профіль"

    days = sorted({x["tenderPeriod"]["startDate"][:10] for x in fake.tenders.values()})
    rep = run(capsys, "coverage", "--from", days[0], "--to", days[-1], "--cpv", "48760000-3,32420000-3")
    assert rep["site_selected"] > 0 and "виключено вручну" in rep["by_reason"]

    # sync by UA-… ids, one of them unknown
    ids = tmp_path / "ids.txt"
    ids.write_text(f"{t['tenderID']}\nUA-2000-01-01-000000-a  # немає\n", encoding="utf-8")
    out = run(capsys, "sync", "--tender-ids", str(ids))
    assert len(out["tenders"]) == 1 and len(out["not_resolved"]) == 1

    assert run(capsys, "exclude", "remove", t["tenderID"])[0]["removed"]
    copy = run(capsys, "export-db", str(tmp_path / "copy.db"), "--relevant-only")
    assert copy["dropped_decisions"] > 0
    assert len(Database(tmp_path / "copy.db").search(limit=None)[0]) == relevant


def test_cli_incomplete_sync_exits_non_zero(served, capsys):
    served.empty_feed_calls = set(range(2, 100))
    with pytest.raises(SystemExit) as e:
        cli.main(["sync", "--since", "2000-01-01"])
    assert e.value.code == 3
    out = capsys.readouterr()
    assert json.loads(out.out)["complete"] is False and "НЕПОВНА" in out.err
