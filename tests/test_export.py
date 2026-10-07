from datetime import datetime

from conftest import NOW
from openpyxl import load_workbook
from test_sync import run_sync

from prozorro_mcp.db import Database
from prozorro_mcp.export import export_tenders
from prozorro_mcp.selection import TenderQuery, select_tenders


async def synced_db(fake, tender_filter) -> Database:
    db = Database(":memory:")
    await run_sync(fake, tender_filter, db)
    return db


async def test_select_by_stage(fake, tender_filter):
    db = await synced_db(fake, tender_filter)
    assert {t["title"] for t in select_tenders(db, TenderQuery(stage="complete"))} == {"Закупівля комутаторів для ЦОД"}
    assert len(select_tenders(db, TenderQuery(stage="active"))) == 2
    assert len(select_tenders(db, TenderQuery(stage="all"))) == 3


async def test_select_by_award_date(fake, tender_filter):
    db = await synced_db(fake, tender_filter)
    day = NOW.strftime("%Y-%m-%d")
    assert len(select_tenders(db, TenderQuery(awarded_from=day))) == 1
    assert select_tenders(db, TenderQuery(awarded_from="2030-01-01")) == []


async def test_export_excel(fake, tender_filter, tmp_path):
    db = await synced_db(fake, tender_filter)
    path = tmp_path / "out.xlsx"
    counts = export_tenders(select_tenders(db, TenderQuery()), path, tender_filter)
    counts.pop("Аналітика")
    assert counts == {"Тендери": 3, "Позиції": 4, "Переможці": 1, "Пропозиції": 2, "Ціни за одиницю": 1}

    wb = load_workbook(path)
    assert wb.sheetnames == ["Аналітика", "Тендери", "Позиції", "Переможці", "Пропозиції", "Ціни за одиницю"]
    summary = {r[0]: r[1] for r in wb["Аналітика"].iter_rows(values_only=True) if r and r[0]}
    assert summary["Тендерів"] == 3
    assert summary["Медіанна знижка від очікуваної вартості"] == 0.18
    ws = wb["Тендери"]
    header = [c.value for c in ws[1]]
    rows = {r[1].value: r for r in ws.iter_rows(min_row=2)}
    switches = rows["Закупівля комутаторів для ЦОД"]
    col = {h: i for i, h in enumerate(header)}
    assert switches[0].hyperlink.target.startswith("https://prozorro.gov.ua/tender/UA-")
    assert switches[col["Очікувана вартість"]].value == 1_200_000
    assert isinstance(switches[col["Створено"]].value, datetime)
    assert switches[col["Переможець"]].value == 'ТОВ "Мережеві Рішення"'
    assert switches[col["Знижка від очікуваної"]].value == 0.18
    assert switches[col["Статус"]].value == "Завершено"
    assert switches[col["Договір"]].value == "Підписано"

    awards = list(wb["Переможці"].iter_rows(min_row=2, values_only=True))
    assert awards[0][4] == "Переможець" and awards[0][6] == "12345678"
    bids = list(wb["Пропозиції"].iter_rows(min_row=2, values_only=True))
    assert sorted(b[8] or "" for b in bids) == ["", "так"]
    prices = list(wb["Ціни за одиницю"].iter_rows(min_row=2, values_only=True))
    assert prices[0][7] == 98_000


async def test_period_mode(fake, tender_filter):
    db = await synced_db(fake, tender_filter)
    day = NOW.strftime("%Y-%m-%d")
    # all three were announced today; only one has a decision today
    assert len(select_tenders(db, TenderQuery(period_from=day, period_mode="created"))) == 3
    assert len(select_tenders(db, TenderQuery(period_from=day, period_mode="awarded"))) == 1
    assert len(select_tenders(db, TenderQuery(period_from=day, period_mode="either"))) == 3
    assert len(select_tenders(db, TenderQuery(period_from="2030-01-01", period_mode="either"))) == 0
